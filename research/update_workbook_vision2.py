"""Second workbook pass: MATRIX graded results, BioReason-Pro, OmniMatBench.

First pass (research/update_workbook_vision.py) covered the MATCHA and CSMBench
Harbor numbers. This pass adds the MATRIX five-level graded reward, the
BioReason-Pro row, and explicit blockers for the rows that cannot be measured.

MATRIX caveat that must not be lost: its reward is a five-level rubric score
(0/0.25/0.5/0.75/1) from a gpt-5.1 judge under a reconstructed official-style
protocol, NOT binary accuracy. It is a proxy metric and cannot support a
qualifying G claim on its own.

Usage:
    python research/update_workbook_vision2.py            # apply
    python research/update_workbook_vision2.py --dry-run  # print diff only
"""

from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "GENESIS Benchmark Brainstorming.xlsx"
BACKUP = ROOT / "research" / "workbook-backup-pre-vision2.xlsx"
RUN_DIR = (
    ROOT / "material-harbor" / "harbor_vision" / "runs" / "matrix-vision-n10"
)

GREEN = PatternFill("solid", fgColor="FFC6EFCE")
YELLOW = PatternFill("solid", fgColor="FFFFEB9C")
RED = PatternFill("solid", fgColor="FFFFC7CE")


def matrix_summary() -> tuple[str, str]:
    """Read the graded MATRIX state from the run directory.

    Reads per-trial result.json files so the workbook can be updated while the
    batch is still running; the aggregate report.json only exists at the end.
    """
    report_path = RUN_DIR / "report.json"
    if report_path.exists():
        report = json.loads(report_path.read_text())
        rewards = report["rewards"]
        scored = [
            r["reward"] for r in rewards.values()
            if isinstance(r.get("reward"), (int, float))
        ]
        total = len(rewards)
    else:
        scored, total = [], 0
        for path in sorted((RUN_DIR / "jobs").glob("*/*/result.json")):
            data = json.loads(path.read_text())
            if not data.get("trial_name"):
                continue
            total += 1
            reward = (data.get("verifier_result") or {}).get("rewards", {}).get("reward")
            if isinstance(reward, (int, float)):
                scored.append(reward)

    mean = (sum(scored) / len(scored)) if scored else 0.0
    g = (
        f"PROXY METRIC, not binary accuracy: mean five-level graded reward "
        f"{mean:.3f} over {len(scored)}/{total} trials "
        f"(Harbor run matrix-vision-n10, deepseek-v4.1-flash, real figures). "
        f"Judge gpt-5.1; protocol reconstructed official-style, NOT official-exact. "
        f"Per-trial levels: {', '.join(f'{v:g}' for v in sorted(scored))}."
    )
    j = (
        "No (G not claimable: the only MATRIX number is a judge-graded five-level "
        "proxy on a reconstructed rubric, not binary accuracy; vision rubrics are "
        "adapter-reconstructed because the release publishes kinds but no rubrics)"
    )
    return g, j


def apply(dry_run: bool) -> dict:
    g_matrix, j_matrix = matrix_summary()
    wb = openpyxl.load_workbook(WORKBOOK)

    edits = [
        ("Sheet1", "G21", g_matrix, YELLOW),
        ("Sheet1", "J21", j_matrix, RED),
        # OmniMatBench: state the concrete blockers rather than a vague pending.
        (
            "Sheet1", "G24",
            "Unmeasured (multimodal). Text-only CAL path measured separately; "
            "no image-bearing Harbor run exists for this row.",
            YELLOW,
        ),
        (
            "Sheet1", "J24",
            "No (not measurable as multimodal here: no image-bearing task set was "
            "staged for Harbor, and the CAL path is text-only, so it cannot "
            "demonstrate image understanding. Needs a staged image-bearing subset.)",
            RED,
        ),
        # BioReason-Pro: pending on input separation, NOT disqualified for a
        # zero-model F1. The official prompt builder uses sequence + organism.
        (
            "Sheet1", "G25",
            "Unmeasured (no Harbor run). The released parquet ships go_ids and "
            "go_pred, but the official prompt builder feeds the model sequence and "
            "organism only, so those are label/baseline columns rather than prompt "
            "leakage as long as an adapter keeps them verifier-side.",
            YELLOW,
        ),
        (
            "Sheet1", "J25",
            "Pending, not disqualified. Requirement is input separation, not a "
            "zero-model-F1 failure: the adapter must strip go_ids/go_pred (and free-text "
            "function/location fields) from solver-visible input, deduplicate against "
            "SFT/GO-GPT training rows, and fix a GO propagation + partial-credit policy "
            "before FP/FN are even definable.",
            YELLOW,
        ),
    ]

    changed = []
    for sheet, coord, value, fill in edits:
        ws = wb[sheet]
        changed.append(
            {"sheet": sheet, "cell": coord, "before": ws[coord].value, "after": value}
        )
        if not dry_run:
            ws[coord] = value
            ws[coord].fill = fill

    if not dry_run:
        wb.save(WORKBOOK)
    return {"changed": changed, "dry_run": dry_run}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    if not args.dry_run and not BACKUP.exists():
        shutil.copyfile(WORKBOOK, BACKUP)
    print(json.dumps(apply(args.dry_run), indent=2))
